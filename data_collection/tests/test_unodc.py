# AI-assisted: written with ChatGPT (OpenAI). See docs/AI_USAGE.md.
"""Regression checks for lost historical years, units, and source-row preservation."""
import datetime as dt
import importlib.util
import sqlite3
from pathlib import Path

import openpyxl
import pytest

spec = importlib.util.spec_from_file_location('unodc_collector', Path(__file__).parents[1] / 'unodc.py')
u = importlib.util.module_from_spec(spec)
spec.loader.exec_module(u)


class Sheet:
    def __init__(self, rows): self.rows = rows
    def to_python(self, **_): return self.rows


class Workbook:
    def __init__(self, sheets): self.sheets = sheets; self.sheet_names = list(sheets)
    def get_sheet_by_name(self, key): return Sheet(self.sheets[key])


def database():
    db = sqlite3.connect(':memory:')
    u.schema(db)
    return db


def test_years_and_missing_are_not_cut_to_1990_or_imputed():
    assert u.year_of(1950) == 1950
    assert u.year_of(dt.date(1980, 1, 1)) == 1980
    assert u.number('..') is None
    assert u.number('') is None
    assert u.number(0) == 0
    assert u.year_of(1990.5) is None


def test_ids_retains_unknown_drug_units_and_zero(tmp_path):
    path = tmp_path / 'ids_test.xlsx'
    w = openpyxl.Workbook(); s = w.active; s.title = '2011'
    s.append(['Seizure Date','Country/Territory of Seizure','ISO3','Drug/Substance','Quantity Seized','Measurement Unit','Future column'])
    s.append([dt.date(2011,1,1),'Austria','AUT','Heroin',2000,'g','kept'])
    s.append([dt.date(2011,1,2),'Austria','AUT','Methamphetamine',200,'Tablets','kept'])
    s.append([dt.date(2011,1,3),'Austria','AUT','Unknown substance',0,'kg','kept'])
    w.save(path)
    db=database(); u.ids(db,path)
    rows=db.execute('SELECT qty,kg,kg_trace_equivalent,extra_json FROM seizures_raw ORDER BY row_no').fetchall()
    assert len(rows)==3
    assert rows[0][:3]==(2000,2,2)
    assert rows[1][1:3]==(None,None)
    assert rows[2][0:2]==(0,0)
    assert all('Future column' in row[3] for row in rows)


def test_us_price_series_keeps_1990_and_distinguishes_adjustments():
    wb=Workbook({'Heroin_US':[
        ['Heroin retail prices, US$ per gram'], ['',1990,1991,1992,1993,1994],
        ['Average, in US$',100,101,102,103,''],
        ['Average, inflation adjusted in 2018 US$',200,201,202,203,204],
        ['Heroin wholesale prices, US$ per kilogram'],
        ['Average, in US$',10000,10100,10200,10300,10400],
    ]})
    db=database();u.price_series(db,'source',wb,{})
    assert db.execute('SELECT COUNT(*) FROM unodc_price_observations').fetchone()[0]==15
    assert db.execute("SELECT value FROM unodc_price_observations WHERE year=1990 AND level='retail' AND basis='nominal'").fetchone()[0]==100
    assert db.execute("SELECT value FROM unodc_price_observations WHERE year=1994 AND level='retail' AND basis='nominal'").fetchone()[0] is None
    assert db.execute("SELECT unit FROM unodc_price_observations WHERE level='wholesale' LIMIT 1").fetchone()[0]=='USD/kg'


def test_cultivation_retains_bounds_and_pre_1990_years():
    wb=Workbook({'Sheet1':[
        ['Country',1980,1981,1982,1983,1984],
        ['Afghanistan (best estimate)',100,110,120,130,'..'],
        ['lower bound a',90,100,110,120,'..'],
        ['upper bound a',110,120,130,140,'..'],
    ]})
    db=database();u.wide_cultivation(db,'source',wb,'opium_poppy','hectares')
    assert db.execute('SELECT MIN(year),COUNT(*) FROM unodc_cultivation_observations').fetchone()==(1980,15)
    assert db.execute("SELECT iso3,value FROM unodc_cultivation_observations WHERE year=1980 AND estimate_type='lower'").fetchone()==('AFG',90)


@pytest.mark.skipif(not (u.CACHE / 'hist2015_opium_pdf.pdf').exists(), reason='UNODC PDF cache unavailable')
def test_archived_pdf_cultivation_keeps_1999_and_page_notes():
    db=database()
    u.pdf_cultivation(db,'hist2015_opium_pdf',u.CACHE / 'hist2015_opium_pdf.pdf')
    assert db.execute("SELECT value FROM unodc_cultivation_observations WHERE iso3='AFG' AND year=1999 AND metric='hectares'").fetchone()==(90583,)
    assert db.execute("SELECT value FROM unodc_cultivation_observations WHERE iso3='AFG' AND year=2014 AND metric='production_t'").fetchone()==(6400,)
    assert db.execute("SELECT value FROM unodc_cultivation_observations WHERE iso3='PAK' AND year=2014 AND metric='production_t'").fetchone()==(None,)
    assert 'not comparable' in db.execute('SELECT text FROM unodc_pdf_pages WHERE page_no=1').fetchone()[0]


@pytest.mark.skipif(not (u.CACHE / 'hist2015_coca_pdf.pdf').exists(), reason='UNODC PDF cache unavailable')
def test_archived_pdf_coca_retains_alternate_peru_concept():
    db=database()
    u.pdf_cultivation(db,'hist2015_coca_pdf',u.CACHE / 'hist2015_coca_pdf.pdf')
    assert db.execute("SELECT value,estimate_type FROM unodc_cultivation_observations WHERE iso3='PER' AND year=2011 ORDER BY value").fetchall()==[
        (62500,'central'),(64400,'alternate_concept')]
